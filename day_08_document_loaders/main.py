"""
Day 08: Document Loaders (PDF, HTML, DOCX, Tables)
Core Concept: Extracting clean text and tables from messy files
Chapter: Search & Context

This script executes the key algorithmic implementations and examples from Day 8.
"""
from __future__ import annotations
import os
import sys

def main():
    print("=" * 60)
    print("Executing Day 08: Document Loaders (PDF, HTML, DOCX, Tables)")
    print("=" * 60)

# --- Snippet 1 ---
chars_per_page = len(page.extract_text().strip())
if chars_per_page < 50:
    # probably scanned (or an image-only page) -> OCR

# --- Snippet 2 ---
from pathlib import Path
import pandas as pd
from docx import Document as DocxDocument
from fpdf import FPDF

OUT = Path("day8_docs")
OUT.mkdir(exist_ok=True)

# --------- PDF: 2 pages, header/footer, hyphenated line break, a table ---------
class GuidePDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 8, "ZyntraCorp Bank - Product Guide 2026 - CONFIDENTIAL", align="C")
        self.ln(12)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", size=8)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")

pdf = GuidePDF()
pdf.add_page()
pdf.set_font("Helvetica", "B", 14)
pdf.cell(0, 10, "1. Home Loans", new_x="LMARGIN", new_y="NEXT")
pdf.set_font("Helvetica", size=11)
pdf.multi_cell(0, 6, "Home loan interest rates start at 7.9% per annum for salaried customers and "
                     "8.3% for self-employed customers. All loans are subject to credit assess-\n"
                     "ment by the bank. The maximum tenure is 30 years.")
pdf.ln(4)
pdf.multi_cell(0, 6, "Processing fees are 0.5% of the loan amount, capped at 10,000 rupees.")

pdf.add_page()
pdf.set_font("Helvetica", "B", 14)
pdf.cell(0, 10, "2. Fixed Deposits", new_x="LMARGIN", new_y="NEXT")
pdf.set_font("Helvetica", size=11)
pdf.multi_cell(0, 6, "The table below lists fixed deposit rates by customer type and tenure.")
pdf.ln(3)
fd_rows = [("Customer type", "Tenure", "Rate"),
           ("Regular", "1-2 years", "7.1%"),
           ("Senior citizen", "1-2 years", "7.4%"),
           ("Regular", "3-5 years", "6.9%")]
with pdf.table() as table:
    for row in fd_rows:
        r = table.row()
        for cell in row:
            r.cell(cell)
pdf.ln(4)
pdf.multi_cell(0, 6, "Premature withdrawal attracts a penalty of 1% on the applicable rate.")
pdf.output(str(OUT / "product_guide_2026.pdf"))

# --------- DOCX: headings, paragraphs, a table (in a meaningful order) ---------
doc = DocxDocument()
doc.add_heading("Credit Card Policy 2026", level=1)
doc.add_heading("Platinum Card", level=2)
doc.add_paragraph("The Platinum credit card has an annual fee of 2,500 rupees, "
                  "waived on annual spends above 3 lakh rupees.")
t = doc.add_table(rows=1, cols=2)
t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Benefit", "Details"
for k, v in [("Cashback", "2% on online purchases"), ("Lounge access", "4 visits per year")]:
    cells = t.add_row().cells
    cells[0].text, cells[1].text = k, v
doc.add_heading("Lost or Stolen Cards", level=2)
doc.add_paragraph("Report a lost card immediately in the mobile app. The card is blocked instantly.")
doc.save(OUT / "card_policy_2026.docx")

# --------- HTML: real-world boilerplate + main article + table ---------
(OUT / "savings_page.html").write_text("""<!DOCTYPE html>
<html><head><title>Savings Plus | ZyntraCorp</title>
<style>.x{color:red}</style><script>trackUser('abc');</script></head>
<body>
<div class="cookie-banner">We use cookies. Accept cookies to continue.</div>
<header><nav><a href="/">Home</a> | <a href="/loans">Loans</a> | <a href="/cards">Cards</a> | <a href="/login">Login</a></nav></header>
<main>
  <h1>Savings Plus Account</h1>
  <p>The Savings Plus account offers 4.25% annual interest, credited quarterly.</p>
  <h2>Eligibility</h2>
  <p>Any resident individual aged 18 or above can open an account. There is no minimum balance requirement.</p>
  <h2>Charges</h2>
  <table>
    <tr><th>Service</th><th>Charge</th></tr>
    <tr><td>UPI transfers</td><td>Free, unlimited</td></tr>
    <tr><td>Debit card (first year)</td><td>Free</td></tr>
    <tr><td>Cheque book (after 25 leaves)</td><td>2 rupees per leaf</td></tr>
  </table>
</main>
<aside>Related: Try our Platinum credit card!</aside>
<footer>(C) 2026 ZyntraCorp Bank | Privacy | Terms | Careers | Contact</footer>
</body></html>""", encoding="utf-8")

# --------- XLSX: a rate sheet ---------
pd.DataFrame({
    "Product": ["Home loan", "Home loan", "Car loan", "Personal loan"],
    "Segment": ["Salaried", "Self-employed", "All", "Salaried"],
    "Rate (%)": [7.9, 8.3, 9.1, 11.5],
    "Max tenure (years)": [30, 30, 7, 5],
}).to_excel(OUT / "loan_rates_2026.xlsx", sheet_name="Rates", index=False)

print("Created:", sorted(p.name for p in OUT.iterdir()))

# --- Snippet 3 ---
from pypdf import PdfReader

reader = PdfReader(OUT / "product_guide_2026.pdf")
pages = [p.extract_text() or "" for p in reader.pages]

for i, text in enumerate(pages, start=1):
    print(f"\n----- PAGE {i} ({len(text)} chars) -----\n{text}")

# Scanned-page check: digital pages have plenty of characters
print("\nLikely scanned pages:", [i for i, t in enumerate(pages, 1) if len(t.strip()) < 50])

# --- Snippet 4 ---
import re
import unicodedata
from collections import Counter

def find_repeated_lines(pages: list[str], min_share: float = 0.5) -> set[str]:
    """Lines (with digits masked) that appear on at least 'min_share' of pages."""
    counts = Counter()
    for text in pages:
        masked = {re.sub(r"\d+", "#", ln.strip()) for ln in text.splitlines() if ln.strip()}
        counts.update(masked)
    threshold = max(2, int(len(pages) * min_share))
    return {line for line, n in counts.items() if n >= threshold}

def clean_page(text: str, repeated: set[str]) -> str:
    text = unicodedata.normalize("NFKC", text)                                # ligatures, odd spaces
    lines = [ln for ln in text.splitlines()
             if re.sub(r"\d+", "#", ln.strip()) not in repeated]              # drop headers/footers
    text = "\n".join(lines)
    text = re.sub(r"(\w+)-\n(\w+)", r"\1\2", text)                           # de-hyphenate line breaks
    text = re.sub(r"[ \t]+", " ", text)                                       # collapse spaces
    text = re.sub(r"\n{3,}", "\n\n", text)                                    # max one blank line
    return text.strip()

repeated = find_repeated_lines(pages)
print("Repeated lines removed:", repeated)
clean_pages = [clean_page(t, repeated) for t in pages]
for i, t in enumerate(clean_pages, 1):
    print(f"\n===== CLEAN PAGE {i} =====\n{t}")

# NFKC in action (ligature + non-breaking space + full-width digits)
print(unicodedata.normalize("NFKC", "\ufb01xed\u00a0deposit rate \uff17.\uff14%"))  # -> "fixed deposit rate 7.4%"

# --- Snippet 5 ---
import pdfplumber

def rows_to_markdown(rows: list[list]) -> str:
    rows = [[(c or "").strip().replace("\n", " ") for c in r] for r in rows if r and any(r)]
    if not rows:
        return ""
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---" for _ in header]) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)

def rows_to_sentences(rows: list[list], title: str = "") -> list[str]:
    header, *body = [[(c or "").strip() for c in r] for r in rows]
    prefix = f"{title}: " if title else ""
    return [prefix + "; ".join(f"{h}: {v}" for h, v in zip(header, r)) for r in body]

with pdfplumber.open(OUT / "product_guide_2026.pdf") as pdf:
    for page_no, page in enumerate(pdf.pages, start=1):
        for table in page.extract_tables():
            print(f"\nTable found on page {page_no}:\n")
            print(rows_to_markdown(table))
            print("\nAs row sentences:")
            for s in rows_to_sentences(table, title="Fixed deposit rates"):
                print(" -", s)

# --- Snippet 6 ---
from bs4 import BeautifulSoup

html = (OUT / "savings_page.html").read_text(encoding="utf-8")

# Naive: everything, including boilerplate
naive = BeautifulSoup(html, "html.parser").get_text(" ", strip=True)
print("NAIVE:", naive[:300], "...\n")

# Clean: remove boilerplate, keep structure
soup = BeautifulSoup(html, "html.parser")
for tag in soup(["script", "style", "nav", "header", "footer", "aside", "form", "noscript"]):
    tag.decompose()
for tag in soup.select('[class*="cookie"], [id*="cookie"]'):
    tag.decompose()
main = soup.find("main") or soup.body

sections, path, buffer = [], [], []

def flush():
    if buffer:
        sections.append({"section": " > ".join(path), "text": "\n".join(buffer)})
        buffer.clear()

for el in main.find_all(["h1", "h2", "h3", "p", "li", "table"]):
    if el.name in ("h1", "h2", "h3"):
        flush()
        level = int(el.name[1])
        path = path[: level - 1] + [el.get_text(strip=True)]
    elif el.name == "table":
        rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])] for tr in el.find_all("tr")]
        buffer.append(rows_to_markdown(rows))
    elif not el.find_parent("table"):
        buffer.append(el.get_text(" ", strip=True))
flush()

title = soup.title.get_text(strip=True) if soup.title else ""
for s in sections:
    print(f"[{title} | {s['section']}]\n{s['text']}\n")

# --- Snippet 7 ---
from docx.table import Table
from docx.text.paragraph import Paragraph

docx_file = DocxDocument(OUT / "card_policy_2026.docx")

# ❌ The common mistake: paragraphs and tables read separately, so their order is lost
print("Paragraphs only:", [p.text for p in docx_file.paragraphs if p.text][:4], "...")
print("Tables only     :", [[c.text for c in r.cells] for r in docx_file.tables[0].rows], "\n")

# ✅ Correct: iterate blocks in document order and track headings
sections, path, buffer = [], [], []

def flush_docx():
    if buffer:
        sections.append({"section": " > ".join(path), "text": "\n".join(buffer)})
        buffer.clear()

for block in docx_file.iter_inner_content():                                   # python-docx >= 1.0
    if isinstance(block, Paragraph):
        style = block.style.name if block.style is not None else ""
        if style.startswith("Heading") and style.split()[-1].isdigit():
            flush_docx()
            level = int(style.split()[-1])
            path = path[: level - 1] + [block.text.strip()]
        elif block.text.strip():
            buffer.append(block.text.strip())
    elif isinstance(block, Table):
        buffer.append(rows_to_markdown([[c.text for c in row.cells] for row in block.rows]))
flush_docx()

for s in sections:
    print(f"[{s['section']}]\n{s['text']}\n")

# --- Snippet 8 ---
xlsx = pd.read_excel(OUT / "loan_rates_2026.xlsx", sheet_name=None)            # dict of DataFrames (all sheets)

for sheet, df in xlsx.items():
    rows = [list(df.columns)] + df.astype(str).values.tolist()
    print(f"--- Sheet '{sheet}' as Markdown (1 chunk) ---")
    print(rows_to_markdown(rows))
    print(f"\n--- Sheet '{sheet}' as row sentences (1 chunk per row) ---")
    for s in rows_to_sentences(rows, title=f"Loan rates 2026 ({sheet})"):
        print(" -", s)

# --- Snippet 9 ---
import hashlib
from dataclasses import dataclass, field

@dataclass
class Doc:
    text: str
    metadata: dict = field(default_factory=dict)

def _meta(path: Path, doc_type: str, parser: str, **extra) -> dict:
    return {"source": path.name, "doc_type": doc_type, "parser": parser, **extra}

def load_pdf(path: Path) -> list[Doc]:
    raw_pages, tables_per_page = [], []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            tables = page.find_tables()
            bboxes = [t.bbox for t in tables]

            def outside_tables(obj, bboxes=bboxes):
                # keep only characters that are NOT inside a detected table (avoids duplicating table text)
                return not any(x0 <= obj.get("x0", -1) and obj.get("x1", -1) <= x1 and
                               top <= obj.get("top", -1) and obj.get("bottom", -1) <= bottom
                               for x0, top, x1, bottom in bboxes)

            raw_pages.append(page.filter(outside_tables).extract_text() or "")
            tables_per_page.append([t.extract() for t in tables])
    repeated = find_repeated_lines(raw_pages)
    docs = []
    for i, (raw, tables) in enumerate(zip(raw_pages, tables_per_page), start=1):
        text = clean_page(raw, repeated)
        for tbl in tables:                                                     # append tables as Markdown
            text += "\n\n" + rows_to_markdown(tbl)
        docs.append(Doc(text, _meta(path, "pdf", "pdfplumber", page=i,
                                    ocr_needed=len(raw.strip()) < 50 and not tables)))
    return docs

def load_html(path: Path) -> list[Doc]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else path.stem
    for tag in soup(["script", "style", "nav", "header", "footer", "aside", "form",
                     "noscript"]):
        tag.decompose()
    for tag in soup.select('[class*="cookie"], [id*="cookie"]'):
        tag.decompose()
    main = soup.find("main") or soup.body
    out, path_, buf = [], [], []

    def flush():
        if buf:
            out.append(Doc("\n".join(buf), _meta(path, "html", "bs4", title=title,
                                                 section=" > ".join(path_))))
            buf.clear()

    for el in main.find_all(["h1", "h2", "h3", "p", "li", "table"]):
        if el.name in ("h1", "h2", "h3"):
            flush()
            path_ = path_[: int(el.name[1]) - 1] + [el.get_text(strip=True)]
        elif el.name == "table":
            buf.append(rows_to_markdown([[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
                                         for tr in el.find_all("tr")]))
        elif not el.find_parent("table"):
            buf.append(el.get_text(" ", strip=True))
    flush()
    return out

def load_docx(path: Path) -> list[Doc]:
    d = DocxDocument(path)
    out, path_, buf = [], [], []

    def flush():
        if buf:
            out.append(Doc("\n".join(buf), _meta(path, "docx", "python-docx", section=" > ".join(path_))))
            buf.clear()

    for block in d.iter_inner_content():
        if isinstance(block, Paragraph):
            style = block.style.name if block.style is not None else ""
            if style.startswith("Heading") and style.split()[-1].isdigit():
                flush()
                path_ = path_[: int(style.split()[-1]) - 1] + [block.text.strip()]
            elif block.text.strip():
                buf.append(block.text.strip())
        elif isinstance(block, Table):
            buf.append(rows_to_markdown([[c.text for c in r.cells] for r in block.rows]))
    flush()
    return out

def load_xlsx(path: Path) -> list[Doc]:
    out = []
    for sheet, df in pd.read_excel(path, sheet_name=None).items():
        rows = [list(df.columns)] + df.astype(str).values.tolist()
        out.append(Doc(rows_to_markdown(rows), _meta(path, "xlsx", "pandas", sheet=sheet)))
    return out

LOADERS = {".pdf": load_pdf, ".html": load_html, ".htm": load_html,
           ".docx": load_docx, ".xlsx": load_xlsx,
           ".txt": lambda p: [Doc(p.read_text(encoding="utf-8"), _meta(p, "txt", "plain"))]}

def load_any(path: Path) -> list[Doc]:
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"No loader for {path.suffix}")
    docs = loader(path)
    for d in docs:
        d.metadata["content_hash"] = hashlib.sha256(d.text.encode("utf-8")).hexdigest()[:16]
    return docs

# ---------- Ingest everything + report ----------
all_docs = []
print(f"{'file':<28} {'type':<5} {'units':>5} {'chars':>6} flags")
for p in sorted(OUT.iterdir()):
    docs = load_any(p)
    all_docs.extend(docs)
    flags = "OCR NEEDED" if any(d.metadata.get("ocr_needed") for d in docs) else ""
    print(f"{p.name:<28} {docs[0].metadata['doc_type']:<5} {len(docs):>5} {sum(len(d.text) for d in docs):>6} {flags}")

# ---------- Fact-survival test ----------
EXPECTED = {
    "product_guide_2026.pdf": ["7.9%", "8.3%", "assessment", "7.4%", "Senior citizen", "10,000 rupees"],
    "card_policy_2026.docx": ["2,500 rupees", "Lounge access", "blocked instantly"],
    "savings_page.html": ["4.25%", "no minimum balance", "2 rupees per leaf"],
    "loan_rates_2026.xlsx": ["Car loan", "9.1", "11.5"],
}
FORBIDDEN = ["CONFIDENTIAL", "Accept cookies", "Page 1", "Careers", "trackUser", "Platinum credit card!"]

text_by_source = {}
for d in all_docs:
    text_by_source[d.metadata["source"]] = text_by_source.get(d.metadata["source"], "") + "\n" + d.text

print("\nFact-survival test:")
for src, facts in EXPECTED.items():
    missing = [f for f in facts if f not in text_by_source.get(src, "")]
    noise = [f for f in FORBIDDEN if f in text_by_source.get(src, "")]
    status = "PASS" if not missing and not noise else "FAIL"
    print(f"{status} {src:<26} missing={missing} noise={noise}")

print(f"\nExample metadata:\n{all_docs[-1].metadata}")

# --- Snippet 10 ---
import pymupdf4llm     # AGPL-licensed PyMuPDF underneath; check before production use

md = pymupdf4llm.to_markdown(str(OUT / "product_guide_2026.pdf"))
print(md[:1500])

# --- Snippet 11 ---
# pip install pytesseract pymupdf pillow     (and install Tesseract itself)
import io
import pymupdf
import pytesseract
from PIL import Image

pdf_doc = pymupdf.open(OUT / "product_guide_2026.pdf")
for page in pdf_doc:
    pix = page.get_pixmap(dpi=300)                                            # render the page as an image
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    text = pytesseract.image_to_string(img, lang="eng")                        # e.g., lang="eng+hin" for Hindi
    print(f"--- OCR page {page.number + 1} ---\n{text[:300]}")

if __name__ == "__main__":
    print("Starting Day 08 script execution...")
